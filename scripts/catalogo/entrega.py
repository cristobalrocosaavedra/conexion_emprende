#!/usr/bin/env python3
"""Catálogo para la APK de pruebas y checklist de ida/vuelta, sin publicar producción."""

import argparse
import json
from pathlib import Path
import sys
from datetime import date

from integrar import (
    ROOT,
    AREA,
    DATE,
    REGIONS,
    baseline,
    read_rows,
    simulate,
    add_assertion,
    atomic_json,
    slug,
    digest,
    json_bytes,
)

APP = ROOT / "content/catalogo/maestro-app.json"
BOOK = ROOT / "docs/catalogo-maestro/Checklist-validacion-Territorio-Vivo.xlsx"
FIXED = [
    "ID interno",
    "Tipo",
    "Región",
    "Provincia",
    "Comuna",
    "Nombre actual",
    "Categorías",
    "Latitud actual",
    "Longitud actual",
    "Fuente actual",
    "URL actual",
    "IDs del maestro",
    "Posible duplicado",
]
EDITABLE = [
    "Resultado",
    "Nombre correcto",
    "Territorio correcto",
    "Coordenadas comprobadas",
    "Fuente revisada",
    "Acceso revisado",
    "Licencia revisada",
    "Nombre corregido",
    "Latitud corregida",
    "Longitud corregida",
    "ID duplicado confirmado",
    "Observaciones",
    "Revisor",
    "Fecha revisión",
    "URL evidencia",
]
HEADERS = FIXED + EDITABLE
KINDS = {
    "territory": "Territorio",
    "attraction": "Atractivo",
    "business": "Emprendimiento",
    "route": "Ruta",
}


def compile_catalog():
    rows = read_rows(AREA / "entrada/territorio_vivo_catalogo_maestro_consolidado.csv")
    state, audit = simulate(rows, baseline())
    entities = {e["id"]: e for e in state["entities"]}
    associations = {}
    duplicates = {}
    for proposal in state["proposals"]:
        matches = proposal["matches"]
        candidate = proposal["candidate"]
        # Asociación reversible de evidencia, sin reemplazar un solo valor previo.
        # Solo nombre EXACTO + comuna + tipo + categoría; nunca por parecido.
        if (
            len(matches) == 1
            and matches[0].get("reason") == "same-name-territory-category"
            and not matches[0].get("coordinateConflict")
        ):
            target = entities[matches[0]["entityId"]]
            for assertion in candidate["assertions"]:
                add_assertion(target, assertion)
            for ref in candidate["externalRefs"]:
                if ref not in target["externalRefs"]:
                    target["externalRefs"].append(ref)
            associations[proposal["id"]] = target["id"]
        else:
            entities[candidate["id"]] = candidate
            duplicates[candidate["id"]] = [m["entityId"] for m in matches]
            associations[proposal["id"]] = candidate["id"]
    for e in entities.values():
        for ref in e["externalRefs"]:
            if ref["source"] == "maestro-entregado":
                associations[ref["id"]] = e["id"]
    if len(associations) != len(rows):
        raise ValueError("Faltan filas del maestro por representar.")
    territories = []
    for e in entities.values():
        if e["type"] != "territory":
            continue
        region_ref = next(
            (t for t in e["territoryIds"] if t.startswith("territory:region:")), e["id"]
        )
        # La primera pertenencia es región en los legados; nodos derivados apuntan al padre.
        region = region_ref.split(":")[2]
        province = (
            e["id"].split(":province:")[1].split(":")[0]
            if ":province:" in e["id"]
            else (
                e["territoryIds"][1].split(":province:")[1]
                if len(e["territoryIds"]) > 1
                else None
            )
        )
        commune = (
            slug(e["name"])
            if e["details"]["level"] == "commune"
            else (
                (e["territoryIds"][-1].split(":commune:")[-1])
                if ":commune:" in e["territoryIds"][-1]
                else None
            )
            if e["territoryIds"]
            else None
        )
        territories.append(
            {
                "id": e["id"],
                "level": e["details"]["level"],
                "name": e["name"],
                "region": region,
                "province": province,
                "commune": commune,
                "parentId": e["details"]["parentId"],
                "cut": None,
                "verification": "pending",
            }
        )
    additions, locality_additions, evidence = [], [], {}
    for e in entities.values():
        originals = [
            a["raw"] for a in e["assertions"] if a["source"] == "maestro-entregado"
        ]
        if not originals:
            continue
        ids = [r["id"] for r in originals]
        legacy = e.get("legacy")
        runtime_id = legacy["id"] if legacy else e["id"]
        if e["type"] in ("attraction", "business"):
            row = originals[0]
            evidence[runtime_id] = {
                "rowIds": ids,
                "sourceLabel": row["fuente"],
                "category": row["subcategoria"],
                "declaredAccess": row["acceso"],
                "possibleDuplicateIds": [
                    entities[i].get("legacy", {}).get("id", i)
                    if entities[i].get("legacy")
                    else i
                    for i in duplicates.get(e["id"], [])
                ],
            }
            if not legacy:
                additions.append(
                    {
                        "id": runtime_id,
                        "region": REGIONS[row["codigo_region"]],
                        "province": slug(row["provincia"]),
                        "commune": slug(row["comuna"]),
                        "locality": row["localidad"] or None,
                        "category": e["categoryIds"][0],
                        "name": e["name"],
                        "summary": "",
                        "description": row["notas"],
                        "latitude": e["location"]["latitude"]
                        if e["location"]
                        else None,
                        "longitude": e["location"]["longitude"]
                        if e["location"]
                        else None,
                        "source": {"url": row["url_fuente"], "label": row["fuente"]}
                        if row["url_fuente"]
                        else None,
                        "verification": "pending",
                        "origin": "master",
                        "importedAt": DATE,
                    }
                )
        elif (
            e["type"] == "territory"
            and e["details"]["level"] == "locality"
            and not legacy
        ):
            row = originals[0]
            locality_additions.append(
                {
                    "id": runtime_id,
                    "name": e["name"],
                    "region": REGIONS[row["codigo_region"]],
                    "commune": slug(row["comuna"]),
                    "latitude": e["location"]["latitude"] if e["location"] else None,
                    "longitude": e["location"]["longitude"] if e["location"] else None,
                }
            )
    app = {
        "schemaVersion": 1,
        "snapshotAt": DATE,
        "scope": "apk-de-pruebas-pendiente-de-validacion",
        "entityCount": len(entities),
        "places": additions,
        "localities": locality_additions,
        "territories": territories,
        "evidence": evidence,
        "records": [{"entityId": associations[row["id"]], "raw": row} for row in rows],
        "associations": [
            {
                "rowId": p["id"],
                "entityId": associations[p["id"]],
                "exact": p["candidate"]["id"] not in duplicates,
                "matches": p["matches"],
            }
            for p in state["proposals"]
        ],
    }
    return app, list(entities.values())


def checklist_rows(app, entities):
    territories = {t["id"]: t for t in app["territories"]}
    region_names = {
        t["region"]: t["name"] for t in app["territories"] if t["level"] == "region"
    }
    rows = []
    for e in entities:
        refs = [territories[t] for t in e["territoryIds"] if t in territories]
        if e["type"] == "territory" and all(t["id"] != e["id"] for t in refs):
            refs.append(territories[e["id"]])
        region = ", ".join(dict.fromkeys(region_names[t["region"]] for t in refs))
        province = ", ".join(t["name"] for t in refs if t["level"] == "province")
        commune = ", ".join(t["name"] for t in refs if t["level"] == "commune")
        source = e["assertions"][0]["raw"]
        label = (
            source.get("fuente")
            or (source.get("source") or {}).get("label")
            or e["assertions"][0]["source"]
        )
        url = source.get("url_fuente") or (source.get("source") or {}).get("url") or ""
        ids = ", ".join(
            r["id"] for r in e["externalRefs"] if r["source"] == "maestro-entregado"
        )
        possible = [
            m["entityId"]
            for a in app["associations"]
            if not a["exact"] and a["entityId"] == e["id"]
            for m in a["matches"]
        ]
        location = e["location"] or {}
        fixed = [
            e["id"],
            KINDS.get(e["type"], e["type"])
            + (": " + e["details"]["level"] if e["type"] == "territory" else ""),
            region,
            province,
            commune,
            e["name"],
            ", ".join(e["categoryIds"]),
            str(location["latitude"]) if "latitude" in location else "",
            str(location["longitude"]) if "longitude" in location else "",
            label,
            url,
            ids,
            ", ".join(possible),
        ]
        rows.append(fixed + ["Pendiente"] * 7 + [""] * (len(EDITABLE) - 7))
    return sorted(rows, key=lambda r: (r[2], r[3], r[4], r[1], r[5], r[0]))


def safe_cell(cell, value):
    # Texto literal: una fuente no debe transformarse en fórmula de Excel.
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def write_workbook(path, app, entities):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    intro = wb.active
    intro.title = "Cómo usar"
    instructions = [
        ["Territorio Vivo — checklist de validación"],
        [
            "1. Abre Checklist y filtra por tu región o comuna. Hay 917 entidades, no 917 atractivos."
        ],
        [
            "2. Conserva las columnas grises (identidad/datos actuales). Completa las columnas verdes."
        ],
        [
            "3. Marca Resultado: Validado, No validado o Pendiente. No borres las filas pendientes."
        ],
        [
            "4. Usa Sí / No / Pendiente / No aplica para cada comprobación. Desconocido sigue Pendiente."
        ],
        [
            "5. Escribe correcciones en sus columnas; no cambies los datos actuales ni inventes accesos."
        ],
        [
            "6. Para Validado: nombre/territorio/fuente Sí; coordenadas Sí si hay punto; revisor, fecha y evidencia."
        ],
        [
            "7. Fecha revisión: AAAA-MM-DD. La evidencia es una URL HTTPS. Un No contradice Validado."
        ],
        [
            "8. Coincidencias muestra 22 asociaciones exactas de evidencia y 2 parecidos sin fusionar."
        ],
        [
            "9. Para duplicados confirmados indica el ID destino existente; no se unen automáticamente."
        ],
        [
            "10. Guarda este .xlsx y devuélvelo completo aquí; se preparará un informe de cambios antes de aplicarlos."
        ],
        [
            "Las coordenadas suministradas son referencias sin corroborar; no son puertas de acceso."
        ],
        [
            "No hay inventario oficial SERNATUR/MMA descargado ni contraste CUT oficial completado."
        ],
        [
            "La validación humana de datos no resuelve por sí sola licencia ni autoriza publicación productiva."
        ],
        [
            "Todas las filas parten Pendiente. Se puede devolver una revisión parcial sin esperar a completar todo."
        ],
    ]
    for row in instructions:
        intro.append(row)
    intro.column_dimensions["A"].width = 145
    for row in intro:
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        intro.row_dimensions[row[0].row].height = 32
    intro["A1"].font = Font(bold=True, size=18, color="087F8C")
    checklist = wb.create_sheet("Checklist")
    checklist.append(HEADERS)
    for values in checklist_rows(app, entities):
        for i, value in enumerate(values, 1):
            safe_cell(
                checklist.cell(checklist.max_row + 1, i)
                if i == 1
                else checklist.cell(checklist.max_row, i),
                value,
            )
    for column in range(1, len(HEADERS) + 1):
        checklist.column_dimensions[get_column_letter(column)].width = (
            24 if column > len(FIXED) else 28
        )
        cell = checklist.cell(1, column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid", fgColor="4E6270" if column <= len(FIXED) else "147D65"
        )
    checklist.freeze_panes = "F2"
    table = Table(
        displayName="CatalogoValidacion",
        ref=f"A1:{get_column_letter(len(HEADERS))}{checklist.max_row}",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    checklist.add_table(table)
    for col, choices in [(14, "Validado,No validado,Pendiente")] + [
        (i, "Sí,No,Pendiente,No aplica") for i in range(15, 21)
    ]:
        validation = DataValidation(
            type="list", formula1='"' + choices + '"', allow_blank=False
        )
        validation.errorTitle = "Elige una opción"
        validation.error = "Selecciona un valor del listado."
        validation.showErrorMessage = True
        checklist.add_data_validation(validation)
        validation.add(
            f"{get_column_letter(col)}2:{get_column_letter(col)}{checklist.max_row}"
        )
    original = wb.create_sheet("Origen maestro")
    original.append(["ID destino"] + list(app["records"][0]["raw"]))
    for record in app["records"]:
        original.append([record["entityId"]] + list(record["raw"].values()))
    matches = wb.create_sheet("Coincidencias")
    matches.append(
        [
            "ID maestro",
            "ID destino",
            "Asociación exacta",
            "Nombre de referencia",
            "ID relacionado",
            "Motivo",
        ]
    )
    for a in app["associations"]:
        for match in a["matches"]:
            matches.append(
                [
                    a["rowId"],
                    a["entityId"],
                    "Sí" if a["exact"] else "No",
                    match.get("name"),
                    match["entityId"],
                    match.get("reason"),
                ]
            )
    # Manifiesto de ida/vuelta: detecta cambios accidentales en las columnas de identidad.
    manifest = wb.create_sheet("_control")
    manifest.append(["esquema", "1"])
    manifest.append(["appSha256", digest(json_bytes(app))])
    for row in checklist_rows(app, entities):
        manifest.append([row[0], digest(json_bytes(row[: len(FIXED)]))])
    manifest.sheet_state = "hidden"
    for ws in [original, matches]:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 28
    # No fórmulas, vínculos ejecutables, macros ni conexiones externas.
    for ws in wb:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_return(path, app, entities):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if "Checklist" not in wb.sheetnames:
            raise ValueError("Falta la hoja Checklist.")
        data = list(wb["Checklist"].iter_rows(values_only=True))
        if list(data[0]) != HEADERS:
            raise ValueError("No cambies los encabezados del checklist.")
        expected = {r[0]: r[: len(FIXED)] for r in checklist_rows(app, entities)}
        ids = set()
        result = []
        errors = []
        for index, values in enumerate(data[1:], 2):
            if all(v is None for v in values):
                continue
            row = dict(zip(HEADERS, values))
            id_ = row["ID interno"]
            if id_ not in expected or id_ in ids:
                raise ValueError("ID desconocido o repetido: " + str(id_))
            ids.add(id_)
            # Excel convierte celdas vacías en None; comparar normalizadas, por ID y no por orden.
            if [v if v is not None else "" for v in values[: len(FIXED)]] != [
                v if v is not None else "" for v in expected[id_]
            ]:
                raise ValueError(
                    "Se modificaron datos actuales de "
                    + id_
                    + ". Usa las columnas de corrección."
                )
            if any(isinstance(v, str) and v.startswith("=") for v in values):
                raise ValueError("No se admiten fórmulas en el checklist.")
            if row["Resultado"] not in ("Validado", "No validado", "Pendiente"):
                raise ValueError("Resultado inválido en " + id_)
            for name in EDITABLE[1:7]:
                if row[name] not in ("Sí", "No", "Pendiente", "No aplica"):
                    raise ValueError("Checklist inválido en " + id_ + ": " + name)
            if row["Resultado"] == "Validado":
                mandatory = [
                    "Nombre correcto",
                    "Territorio correcto",
                    "Fuente revisada",
                ]
                if row["Latitud actual"] not in (None, ""):
                    mandatory.append("Coordenadas comprobadas")
                if (
                    any(row[k] != "Sí" for k in mandatory)
                    or any(row[k] == "No" for k in EDITABLE[1:7])
                    or not all(
                        row[k] for k in ["Revisor", "Fecha revisión", "URL evidencia"]
                    )
                ):
                    errors.append(
                        {"id": id_, "reason": "validacion-incompleta-o-contradictoria"}
                    )
            if isinstance(row["Fecha revisión"], date):
                row["Fecha revisión"] = (
                    row["Fecha revisión"].date().isoformat()
                    if hasattr(row["Fecha revisión"], "date")
                    else row["Fecha revisión"].isoformat()
                )
            if row["Fecha revisión"]:
                try:
                    date.fromisoformat(str(row["Fecha revisión"]))
                except ValueError:
                    errors.append(
                        {"id": id_, "reason": "fecha-invalida-usar-AAAA-MM-DD"}
                    )
            if row["URL evidencia"] and not str(row["URL evidencia"]).startswith(
                "https://"
            ):
                errors.append({"id": id_, "reason": "evidencia-no-https"})
            target = row["ID duplicado confirmado"]
            if target and (target not in expected or target == id_):
                errors.append({"id": id_, "reason": "id-duplicado-destino-invalido"})
            corrections = (row["Latitud corregida"], row["Longitud corregida"])
            if any(v is not None and v != "" for v in corrections):
                from integrar import coordinate

                try:
                    coordinate(*corrections)
                except (ValueError, TypeError):
                    errors.append(
                        {"id": id_, "reason": "coordenadas-corregidas-invalidas"}
                    )
            result.append({"id": id_, **{key: row[key] for key in EDITABLE}})
        if ids != set(expected):
            raise ValueError(
                "Faltan filas: devuelve el checklist completo, sin borrar pendientes."
            )
        return {
            "schemaVersion": 1,
            "mode": "revision-humana-no-aplicada",
            "rows": result,
            "counts": {
                status: sum(r["Resultado"] == status for r in result)
                for status in ("Validado", "No validado", "Pendiente")
            },
            "conflicts": errors,
            "publications": 0,
        }
    finally:
        wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["generate", "review"])
    parser.add_argument("--xlsx", type=Path, default=BOOK)
    parser.add_argument(
        "--output", type=Path, default=AREA / "trabajo/revision-humana.json"
    )
    args = parser.parse_args()
    app, entities = compile_catalog()
    if args.command == "generate":
        atomic_json(APP, app)
        write_workbook(args.xlsx, app, entities)
        print(
            f"APK: {len(app['places'])} lugares adicionales; {len(app['records'])} filas trazadas. Checklist: {len(entities)} entidades."
        )
    else:
        result = read_return(args.xlsx, app, entities)
        if not args.output.resolve().is_relative_to((AREA / "trabajo").resolve()):
            parser.error(
                "La revisión solo se guarda en trabajo/, sin modificar el catálogo servido."
            )
        atomic_json(args.output, result)
        print(
            json.dumps(
                {
                    "counts": result["counts"],
                    "conflicts": len(result["conflicts"]),
                    "publications": 0,
                },
                ensure_ascii=False,
            )
        )
        if result["conflicts"]:
            sys.exit(2)


if __name__ == "__main__":
    main()
