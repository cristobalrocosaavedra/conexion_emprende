"""Integración del maestro y vuelta del checklist con correcciones de ejemplo."""

import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from datetime import datetime

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts/catalogo"))
SPEC = importlib.util.spec_from_file_location(
    "entrega", ROOT / "scripts/catalogo/entrega.py"
)
m = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(m)


class EntregaCatalogo(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app, cls.entities = m.compile_catalog()
        cls.temp = tempfile.TemporaryDirectory()
        cls.book = Path(cls.temp.name) / "Checklist-ejemplo.xlsx"
        m.write_workbook(cls.book, cls.app, cls.entities)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def modify(self, changes):
        from openpyxl import load_workbook

        book = load_workbook(self.book)
        sheet = book["Checklist"]
        for row, column, value in changes:
            sheet.cell(row, m.HEADERS.index(column) + 1, value)
        path = Path(self.temp.name) / "Checklist-modificado-ejemplo.xlsx"
        book.save(path)
        return path

    def test_todas_las_filas_trazables_y_sin_reemplazar_lugares(self):
        self.assertEqual(len(self.app["records"]), 737)
        self.assertEqual(len(self.entities), 917)
        self.assertEqual(len(self.app["places"]), 365)
        self.assertEqual(len(self.app["territories"]), 485)
        self.assertEqual(
            sum(t["level"] == "commune" for t in self.app["territories"]), 346
        )
        self.assertEqual(sum(a["exact"] for a in self.app["associations"]), 22)
        self.assertEqual(sum(not a["exact"] for a in self.app["associations"]), 2)
        self.assertTrue(all(p["verification"] == "pending" for p in self.app["places"]))
        existing = m.baseline()
        for e in existing["entities"]:
            actual = next(x for x in self.entities if x["id"] == e["id"])
            for key in ("name", "location", "legacy", "details", "review"):
                self.assertEqual(e[key], actual[key])
        self.assertEqual(json.loads(m.APP.read_text()), self.app)
        self.assertEqual(len({r["entityId"] for r in self.app["records"]}), 737)

    def test_excel_sin_tocar_es_revision_parcial_valida(self):
        result = m.read_return(self.book, self.app, self.entities)
        self.assertEqual(
            result["counts"], {"Validado": 0, "No validado": 0, "Pendiente": 917}
        )
        self.assertEqual(result["conflicts"], [])
        self.assertEqual(result["publications"], 0)
        from openpyxl import load_workbook

        book = load_workbook(self.book, read_only=True)
        self.assertEqual(book["Origen maestro"].max_row, 738)
        self.assertEqual(book["Checklist"].max_row, 918)
        self.assertEqual(book["Coincidencias"].max_row, 25)
        self.assertFalse(
            any(c.data_type == "f" for ws in book for row in ws for c in row)
        )
        book.close()

    def test_resultado_validado_exige_evidencia_y_comprobaciones(self):
        path = self.modify([(2, "Resultado", "Validado")])
        result = m.read_return(path, self.app, self.entities)
        self.assertTrue(result["conflicts"])
        path = self.modify(
            [
                (2, "Resultado", "Validado"),
                (2, "Nombre correcto", "Sí"),
                (2, "Territorio correcto", "Sí"),
                (2, "Coordenadas comprobadas", "Sí"),
                (2, "Fuente revisada", "Sí"),
                (2, "Revisor", "Revisor de ejemplo"),
                (2, "Fecha revisión", datetime(2026, 10, 26)),
                (2, "URL evidencia", "https://example.org/evidencia-de-prueba"),
            ]
        )
        result = m.read_return(path, self.app, self.entities)
        self.assertEqual(result["counts"]["Validado"], 1)
        self.assertEqual(result["counts"]["Pendiente"], 916)
        self.assertEqual(result["conflicts"], [])

    def test_correcciones_no_modifican_el_catalogo_y_coordenadas_invalidas(self):
        before = m.APP.read_bytes()
        path = self.modify(
            [
                (2, "Resultado", "No validado"),
                (2, "Nombre corregido", "Nombre corregido de ejemplo"),
                (2, "Latitud corregida", -999),
                (2, "Longitud corregida", -71),
            ]
        )
        result = m.read_return(path, self.app, self.entities)
        self.assertTrue(result["conflicts"])
        self.assertEqual(result["counts"]["No validado"], 1)
        self.assertEqual(m.APP.read_bytes(), before)

    def test_rechaza_ids_actuales_alterados_y_formulas(self):
        for column, value in [
            ("ID interno", "ID-inexistente-ejemplo"),
            ("Nombre actual", "Cambio indebido de ejemplo"),
            ("Observaciones", "=1+1"),
        ]:
            with self.assertRaises(ValueError):
                m.read_return(
                    self.modify([(2, column, value)]), self.app, self.entities
                )

    def test_devuelve_todo_sin_duplicar_ni_borrar_filas(self):
        from openpyxl import load_workbook

        book = load_workbook(self.book)
        ws = book["Checklist"]
        ws.delete_rows(2)
        path = Path(self.temp.name) / "Checklist-faltante-ejemplo.xlsx"
        book.save(path)
        with self.assertRaises(ValueError):
            m.read_return(path, self.app, self.entities)
        path = self.modify(
            [(3, "ID interno", m.checklist_rows(self.app, self.entities)[0][0])]
        )
        with self.assertRaises(ValueError):
            m.read_return(path, self.app, self.entities)


if __name__ == "__main__":
    unittest.main()
